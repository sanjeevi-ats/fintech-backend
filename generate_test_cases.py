import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import json

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "API Test Cases"

# Define headers
headers = [
    "Test Case ID",
    "Endpoint Name",
    "Test Case Description",
    "Request Payload",
    "Expected Status Code",
    "Expected Response",
    "Status",
    "Comments"
]

# Style definitions
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write headers
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border

# Sample GUIDs for testing
SAMPLE_BRANCH_ID = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
SAMPLE_CUSTOMER_ID = "b2c3d4e5-f6a7-8901-bcde-f12345678901"
SAMPLE_LOAN_ID = "c3d4e5f6-a7b8-9012-cdef-123456789012"
SAMPLE_INSTALLMENT_ID = "d4e5f6a7-b8c9-0123-def1-234567890123"
SAMPLE_PARTNER_ID = "e5f6a7b8-c9d0-1234-ef12-345678901234"
SAMPLE_USER_ID = "f6a7b8c9-d0e1-2345-f123-456789012345"
SAMPLE_PRODUCT_ID = "a7b8c9d0-e1f2-3456-1234-567890123456"

# Test cases data
test_cases = [
    # AUTH CONTROLLER
    {
        "id": "TS_API_001",
        "endpoint": "POST /api/v1/auth/login",
        "description": "Verify whether the login endpoint returns 200 response while hitting it with valid credentials",
        "payload": json.dumps({
            "email": "admin@fintech.com",
            "password": "Admin@123"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "token": "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...",
            "email": "admin@fintech.com",
            "name": "Admin User",
            "role": "super_admin"
        }, indent=2)
    },
    {
        "id": "TS_API_002",
        "endpoint": "POST /api/v1/auth/login",
        "description": "Verify whether the login endpoint returns 401 response when invalid email is provided",
        "payload": json.dumps({
            "email": "invalid@fintech.com",
            "password": "Admin@123"
        }, indent=2),
        "status_code": 401,
        "response": json.dumps({
            "statusCode": 401,
            "errorCode": "UNAUTHORIZED",
            "message": "Invalid email or password"
        }, indent=2)
    },
    {
        "id": "TS_API_003",
        "endpoint": "POST /api/v1/auth/login",
        "description": "Verify whether the login endpoint returns 401 response when invalid password is provided",
        "payload": json.dumps({
            "email": "admin@fintech.com",
            "password": "WrongPassword"
        }, indent=2),
        "status_code": 401,
        "response": json.dumps({
            "statusCode": 401,
            "errorCode": "UNAUTHORIZED",
            "message": "Invalid email or password"
        }, indent=2)
    },
    {
        "id": "TS_API_004",
        "endpoint": "POST /api/v1/auth/login",
        "description": "Verify whether the login endpoint returns 400 response when email field is missing",
        "payload": json.dumps({
            "password": "Admin@123"
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Email": ["The Email field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_005",
        "endpoint": "POST /api/v1/auth/login",
        "description": "Verify whether the login endpoint returns 400 response when password field is missing",
        "payload": json.dumps({
            "email": "admin@fintech.com"
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Password": ["The Password field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_006",
        "endpoint": "POST /api/v1/auth/register",
        "description": "Verify whether the register endpoint returns 200 response with valid registration data",
        "payload": json.dumps({
            "name": "John Doe",
            "email": "john.doe@fintech.com",
            "password": "SecurePass@123",
            "branchId": SAMPLE_BRANCH_ID,
            "role": "agent"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "id": "new-user-guid-here",
            "message": "User registered successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_007",
        "endpoint": "POST /api/v1/auth/register",
        "description": "Verify whether the register endpoint returns 400 response when email format is invalid",
        "payload": json.dumps({
            "name": "John Doe",
            "email": "invalid-email",
            "password": "SecurePass@123",
            "branchId": SAMPLE_BRANCH_ID,
            "role": "agent"
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Email": ["The Email field is not a valid e-mail address."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_008",
        "endpoint": "POST /api/v1/auth/register",
        "description": "Verify whether the register endpoint returns 400 response when password is less than 6 characters",
        "payload": json.dumps({
            "name": "John Doe",
            "email": "john.doe@fintech.com",
            "password": "12345",
            "branchId": SAMPLE_BRANCH_ID,
            "role": "agent"
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Password": ["The field Password must be a string with a minimum length of 6."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_009",
        "endpoint": "POST /api/v1/auth/refresh-token",
        "description": "Verify whether the refresh-token endpoint returns 200 response with valid token",
        "payload": json.dumps({
            "token": "valid-refresh-token-here"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "token": "new-jwt-token-here",
            "refreshToken": "new-refresh-token-here"
        }, indent=2)
    },
    {
        "id": "TS_API_010",
        "endpoint": "POST /api/v1/auth/refresh-token",
        "description": "Verify whether the refresh-token endpoint returns 401 response with invalid token",
        "payload": json.dumps({
            "token": "invalid-token"
        }, indent=2),
        "status_code": 401,
        "response": json.dumps({
            "statusCode": 401,
            "message": "Invalid or expired token"
        }, indent=2)
    },
    {
        "id": "TS_API_011",
        "endpoint": "POST /api/v1/auth/enable-totp",
        "description": "Verify whether the enable-totp endpoint returns 200 response with valid user ID",
        "payload": json.dumps({
            "userId": SAMPLE_USER_ID
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "qrCode": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAA...",
            "secret": "JBSWY3DPEHPK3PXP"
        }, indent=2)
    },
    {
        "id": "TS_API_012",
        "endpoint": "POST /api/v1/auth/enable-totp",
        "description": "Verify whether the enable-totp endpoint returns 404 response with non-existent user ID",
        "payload": json.dumps({
            "userId": "00000000-0000-0000-0000-000000000000"
        }, indent=2),
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "User not found"
        }, indent=2)
    },
    {
        "id": "TS_API_013",
        "endpoint": "POST /api/v1/auth/change-password",
        "description": "Verify whether the change-password endpoint returns 200 response with valid credentials",
        "payload": json.dumps({
            "userId": SAMPLE_USER_ID,
            "oldPassword": "OldPass@123",
            "newPassword": "NewPass@123"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "message": "Password changed successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_014",
        "endpoint": "POST /api/v1/auth/change-password",
        "description": "Verify whether the change-password endpoint returns 401 response when old password is incorrect",
        "payload": json.dumps({
            "userId": SAMPLE_USER_ID,
            "oldPassword": "WrongOldPass",
            "newPassword": "NewPass@123"
        }, indent=2),
        "status_code": 401,
        "response": json.dumps({
            "statusCode": 401,
            "message": "Old password is incorrect"
        }, indent=2)
    },
    
    # BRANCH CONTROLLER
    {
        "id": "TS_API_015",
        "endpoint": "GET /api/v1/branch",
        "description": "Verify whether the get all branches endpoint returns 200 response with list of branches",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_BRANCH_ID,
                "name": "Main Branch",
                "city": "Mumbai",
                "isActive": True
            }
        ], indent=2)
    },
    {
        "id": "TS_API_016",
        "endpoint": f"GET /api/v1/branch/{SAMPLE_BRANCH_ID}",
        "description": "Verify whether the get branch by ID endpoint returns 200 response with branch details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "id": SAMPLE_BRANCH_ID,
            "name": "Main Branch",
            "city": "Mumbai",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_017",
        "endpoint": "GET /api/v1/branch/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get branch by ID endpoint returns 404 response when branch does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Branch not found"
        }, indent=2)
    },
    {
        "id": "TS_API_018",
        "endpoint": "GET /api/v1/branch/search?query=Mumbai",
        "description": "Verify whether the search branches endpoint returns 200 response with matching branches",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_BRANCH_ID,
                "name": "Main Branch",
                "city": "Mumbai",
                "isActive": True
            }
        ], indent=2)
    },
    {
        "id": "TS_API_019",
        "endpoint": "POST /api/v1/branch",
        "description": "Verify whether the create branch endpoint returns 201 response with valid branch data",
        "payload": json.dumps({
            "name": "New Branch",
            "city": "Delhi",
            "isActive": True
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-branch-guid",
            "name": "New Branch",
            "city": "Delhi",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_020",
        "endpoint": "POST /api/v1/branch",
        "description": "Verify whether the create branch endpoint returns 400 response when name is missing",
        "payload": json.dumps({
            "city": "Delhi",
            "isActive": True
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Name": ["The Name field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_021",
        "endpoint": "POST /api/v1/branch",
        "description": "Verify whether the create branch endpoint returns 403 response when user lacks MANAGE_BRANCHES permission",
        "payload": json.dumps({
            "name": "New Branch",
            "city": "Delhi",
            "isActive": True
        }, indent=2),
        "status_code": 403,
        "response": json.dumps({
            "statusCode": 403,
            "message": "Insufficient permissions"
        }, indent=2)
    },
    {
        "id": "TS_API_022",
        "endpoint": f"PUT /api/v1/branch/{SAMPLE_BRANCH_ID}",
        "description": "Verify whether the update branch endpoint returns 204 response with valid branch data",
        "payload": json.dumps({
            "id": SAMPLE_BRANCH_ID,
            "name": "Updated Branch",
            "city": "Mumbai",
            "isActive": True
        }, indent=2),
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_023",
        "endpoint": f"PUT /api/v1/branch/{SAMPLE_BRANCH_ID}",
        "description": "Verify whether the update branch endpoint returns 400 response when ID in URL doesn't match ID in body",
        "payload": json.dumps({
            "id": "different-guid-here",
            "name": "Updated Branch",
            "city": "Mumbai",
            "isActive": True
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "ID mismatch"
        }, indent=2)
    },
    {
        "id": "TS_API_024",
        "endpoint": f"DELETE /api/v1/branch/{SAMPLE_BRANCH_ID}",
        "description": "Verify whether the delete branch endpoint returns 204 response when branch is successfully deleted",
        "payload": "N/A",
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_025",
        "endpoint": "DELETE /api/v1/branch/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the delete branch endpoint returns 404 response when branch does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Branch not found"
        }, indent=2)
    },
    {
        "id": "TS_API_026",
        "endpoint": f"PUT /api/v1/branch/{SAMPLE_BRANCH_ID}/settings",
        "description": "Verify whether the update branch settings endpoint returns 204 response with valid settings JSON",
        "payload": json.dumps({
            "maxLoanAmount": 500000,
            "interestRateRange": {"min": 10, "max": 24}
        }, indent=2),
        "status_code": 204,
        "response": "No Content"
    },
    
    # USERS CONTROLLER
    {
        "id": "TS_API_027",
        "endpoint": "GET /api/v1/users",
        "description": "Verify whether the get all users endpoint returns 200 response with list of users",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_USER_ID,
                "name": "John Doe",
                "email": "john.doe@fintech.com",
                "role": "agent",
                "isActive": True
            }
        ], indent=2)
    },
    {
        "id": "TS_API_028",
        "endpoint": f"GET /api/v1/users/{SAMPLE_USER_ID}",
        "description": "Verify whether the get user by ID endpoint returns 200 response with user details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "id": SAMPLE_USER_ID,
            "name": "John Doe",
            "email": "john.doe@fintech.com",
            "role": "agent",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_029",
        "endpoint": "GET /api/v1/users/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get user by ID endpoint returns 404 response when user does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "User not found"
        }, indent=2)
    },
    {
        "id": "TS_API_030",
        "endpoint": "POST /api/v1/users",
        "description": "Verify whether the create user endpoint returns 201 response with valid user data",
        "payload": json.dumps({
            "name": "Jane Smith",
            "email": "jane.smith@fintech.com",
            "password": "SecurePass@123",
            "branchId": SAMPLE_BRANCH_ID,
            "role": "loan_officer"
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-user-guid",
            "name": "Jane Smith",
            "email": "jane.smith@fintech.com",
            "role": "loan_officer",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_031",
        "endpoint": "POST /api/v1/users",
        "description": "Verify whether the create user endpoint returns 400 response when email is invalid",
        "payload": json.dumps({
            "name": "Jane Smith",
            "email": "invalid-email",
            "password": "SecurePass@123",
            "branchId": SAMPLE_BRANCH_ID,
            "role": "loan_officer"
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Email": ["The Email field is not a valid e-mail address."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_032",
        "endpoint": f"PUT /api/v1/users/{SAMPLE_USER_ID}",
        "description": "Verify whether the update user endpoint returns 204 response with valid user data",
        "payload": json.dumps({
            "id": SAMPLE_USER_ID,
            "name": "John Doe Updated",
            "email": "john.doe@fintech.com",
            "role": "agent",
            "isActive": True
        }, indent=2),
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_033",
        "endpoint": f"DELETE /api/v1/users/{SAMPLE_USER_ID}",
        "description": "Verify whether the delete user endpoint returns 204 response when user is successfully deleted",
        "payload": "N/A",
        "status_code": 204,
        "response": "No Content"
    },
    
    # CUSTOMERS CONTROLLER
    {
        "id": "TS_API_034",
        "endpoint": "GET /api/v1/customers",
        "description": "Verify whether the get all customers endpoint returns 200 response with list of customers",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_CUSTOMER_ID,
                "name": "Rajesh Kumar",
                "aadhaar": "****-****-1234",
                "pan": "ABCDE****F"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_035",
        "endpoint": f"GET /api/v1/customers/{SAMPLE_CUSTOMER_ID}",
        "description": "Verify whether the get customer by ID endpoint returns 200 response with customer details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "id": SAMPLE_CUSTOMER_ID,
            "name": "Rajesh Kumar",
            "aadhaar": "****-****-1234",
            "pan": "ABCDE****F"
        }, indent=2)
    },
    {
        "id": "TS_API_036",
        "endpoint": "GET /api/v1/customers/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get customer by ID endpoint returns 404 response when customer does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Customer not found"
        }, indent=2)
    },
    {
        "id": "TS_API_037",
        "endpoint": "POST /api/v1/customers",
        "description": "Verify whether the create customer endpoint returns 201 response with valid customer data",
        "payload": json.dumps({
            "name": "Priya Sharma",
            "aadhaar": "123456789012",
            "pan": "ABCDE1234F",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-customer-guid",
            "name": "Priya Sharma",
            "aadhaar": "****-****-9012",
            "pan": "ABCDE****F"
        }, indent=2)
    },
    {
        "id": "TS_API_038",
        "endpoint": "POST /api/v1/customers",
        "description": "Verify whether the create customer endpoint returns 400 response when name is missing",
        "payload": json.dumps({
            "aadhaar": "123456789012",
            "pan": "ABCDE1234F",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Name": ["The Name field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_039",
        "endpoint": "POST /api/v1/customers",
        "description": "Verify whether the create customer endpoint returns 400 response when Aadhaar is missing",
        "payload": json.dumps({
            "name": "Priya Sharma",
            "pan": "ABCDE1234F",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Aadhaar": ["The Aadhaar field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_040",
        "endpoint": f"PUT /api/v1/customers/{SAMPLE_CUSTOMER_ID}",
        "description": "Verify whether the update customer endpoint returns 204 response with valid customer data",
        "payload": json.dumps({
            "name": "Rajesh Kumar Updated",
            "aadhaar": "123456789012",
            "pan": "ABCDE1234F",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_041",
        "endpoint": f"DELETE /api/v1/customers/{SAMPLE_CUSTOMER_ID}",
        "description": "Verify whether the delete customer endpoint returns 204 response when customer is successfully deleted",
        "payload": "N/A",
        "status_code": 204,
        "response": "No Content"
    },
]

# Continue with more test cases...

    # LOAN CASES CONTROLLER
    {
        "id": "TS_API_042",
        "endpoint": "GET /api/v1/loancases",
        "description": "Verify whether the get all loan cases endpoint returns 200 response with list of loan cases",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_LOAN_ID,
                "customerId": SAMPLE_CUSTOMER_ID,
                "customerName": "Rajesh Kumar",
                "principal": 10000000,
                "interestAmount": 2000000,
                "totalReceivable": 12000000,
                "processingFees": 100000,
                "status": "active"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_043",
        "endpoint": f"GET /api/v1/loancases/{SAMPLE_LOAN_ID}",
        "description": "Verify whether the get loan case by ID endpoint returns 200 response with loan case details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "id": SAMPLE_LOAN_ID,
            "customerId": SAMPLE_CUSTOMER_ID,
            "customerName": "Rajesh Kumar",
            "principal": 10000000,
            "interestAmount": 2000000,
            "totalReceivable": 12000000,
            "processingFees": 100000,
            "status": "active"
        }, indent=2)
    },
    {
        "id": "TS_API_044",
        "endpoint": "GET /api/v1/loancases/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get loan case by ID endpoint returns 404 response when loan case does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    {
        "id": "TS_API_045",
        "endpoint": "POST /api/v1/loancases",
        "description": "Verify whether the create loan case endpoint returns 201 response with valid loan data",
        "payload": json.dumps({
            "customerId": SAMPLE_CUSTOMER_ID,
            "principal": 10000000,
            "interestAmount": 2000000,
            "processingFees": 100000
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-loan-guid",
            "customerId": SAMPLE_CUSTOMER_ID,
            "customerName": "Rajesh Kumar",
            "principal": 10000000,
            "interestAmount": 2000000,
            "totalReceivable": 12000000,
            "processingFees": 100000,
            "status": "draft"
        }, indent=2)
    },
    {
        "id": "TS_API_046",
        "endpoint": "POST /api/v1/loancases",
        "description": "Verify whether the create loan case endpoint returns 400 response when customerId is missing",
        "payload": json.dumps({
            "principal": 10000000,
            "interestAmount": 2000000,
            "processingFees": 100000
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "CustomerId": ["The CustomerId field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_047",
        "endpoint": "POST /api/v1/loancases",
        "description": "Verify whether the create loan case endpoint returns 400 response when principal is missing",
        "payload": json.dumps({
            "customerId": SAMPLE_CUSTOMER_ID,
            "interestAmount": 2000000,
            "processingFees": 100000
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Principal": ["The Principal field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_048",
        "endpoint": f"POST /api/v1/loancases/{SAMPLE_LOAN_ID}/approve",
        "description": "Verify whether the approve loan endpoint returns 200 response when loan is successfully approved",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "message": "Loan approved successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_049",
        "endpoint": "POST /api/v1/loancases/00000000-0000-0000-0000-000000000000/approve",
        "description": "Verify whether the approve loan endpoint returns 404 response when loan does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    {
        "id": "TS_API_050",
        "endpoint": f"POST /api/v1/loancases/{SAMPLE_LOAN_ID}/disburse",
        "description": "Verify whether the disburse loan endpoint returns 200 response when loan is successfully disbursed",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "message": "Loan disbursed successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_051",
        "endpoint": "POST /api/v1/loancases/00000000-0000-0000-0000-000000000000/disburse",
        "description": "Verify whether the disburse loan endpoint returns 404 response when loan does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    
    # INSTALLMENTS CONTROLLER
    {
        "id": "TS_API_052",
        "endpoint": f"POST /api/v1/installments/generate/{SAMPLE_LOAN_ID}?count=12",
        "description": "Verify whether the generate installments endpoint returns 200 response with valid loan ID",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "message": "Installments generated successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_053",
        "endpoint": "POST /api/v1/installments/generate/00000000-0000-0000-0000-000000000000?count=12",
        "description": "Verify whether the generate installments endpoint returns 404 response when loan does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    {
        "id": "TS_API_054",
        "endpoint": "GET /api/v1/installments/due?from=2026-04-01&to=2026-04-30",
        "description": "Verify whether the get due installments endpoint returns 200 response with installments due in date range",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_INSTALLMENT_ID,
                "no": 1,
                "dueDate": "2026-04-15T00:00:00Z",
                "amount": 1000000,
                "status": "due"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_055",
        "endpoint": "GET /api/v1/installments/due?from=2026-05-01&to=2026-04-01",
        "description": "Verify whether the get due installments endpoint returns 400 response when from date is after to date",
        "payload": "N/A",
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "Invalid date range"
        }, indent=2)
    },
    {
        "id": "TS_API_056",
        "endpoint": f"GET /api/v1/installments/loan/{SAMPLE_LOAN_ID}",
        "description": "Verify whether the get installments by loan ID endpoint returns 200 response with list of installments",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_INSTALLMENT_ID,
                "no": 1,
                "dueDate": "2026-04-15T00:00:00Z",
                "amount": 1000000,
                "status": "due"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_057",
        "endpoint": "GET /api/v1/installments/loan/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get installments by loan ID endpoint returns 404 response when loan does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    
    # COLLECTION CONTROLLER
    {
        "id": "TS_API_058",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 200 response with valid payment data",
        "payload": json.dumps({
            "installmentId": SAMPLE_INSTALLMENT_ID,
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "id": "receipt-guid",
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": "",
            "capturedAt": "2026-04-11T10:30:00Z"
        }, indent=2)
    },
    {
        "id": "TS_API_059",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 200 response with UPI payment mode",
        "payload": json.dumps({
            "installmentId": SAMPLE_INSTALLMENT_ID,
            "amountPaid": 1000000,
            "mode": "upi",
            "utrRef": "UPI123456789"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "id": "receipt-guid",
            "amountPaid": 1000000,
            "mode": "upi",
            "utrRef": "UPI123456789",
            "capturedAt": "2026-04-11T10:30:00Z"
        }, indent=2)
    },
    {
        "id": "TS_API_060",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 200 response with bank transfer payment mode",
        "payload": json.dumps({
            "installmentId": SAMPLE_INSTALLMENT_ID,
            "amountPaid": 1000000,
            "mode": "bank_transfer",
            "utrRef": "NEFT987654321"
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "id": "receipt-guid",
            "amountPaid": 1000000,
            "mode": "bank_transfer",
            "utrRef": "NEFT987654321",
            "capturedAt": "2026-04-11T10:30:00Z"
        }, indent=2)
    },
    {
        "id": "TS_API_061",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 400 response when installmentId is missing",
        "payload": json.dumps({
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "InstallmentId": ["The InstallmentId field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_062",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 400 response when amountPaid is missing",
        "payload": json.dumps({
            "installmentId": SAMPLE_INSTALLMENT_ID,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "AmountPaid": ["The AmountPaid field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_063",
        "endpoint": "POST /api/collection/collect",
        "description": "Verify whether the collect payment endpoint returns 404 response when installment does not exist",
        "payload": json.dumps({
            "installmentId": "00000000-0000-0000-0000-000000000000",
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Installment not found"
        }, indent=2)
    },
    {
        "id": "TS_API_064",
        "endpoint": "POST /api/collection/sync",
        "description": "Verify whether the sync offline collections endpoint returns 200 response with valid collection data",
        "payload": json.dumps([
            {
                "localId": "offline-1",
                "installmentId": SAMPLE_INSTALLMENT_ID,
                "amountPaid": 1000000,
                "mode": "cash",
                "utrRef": ""
            }
        ], indent=2),
        "status_code": 200,
        "response": json.dumps([
            {
                "localId": "offline-1",
                "success": True,
                "error": "",
                "receiptId": "receipt-guid"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_065",
        "endpoint": "POST /api/collection/sync",
        "description": "Verify whether the sync offline collections endpoint returns 200 response with multiple collections",
        "payload": json.dumps([
            {
                "localId": "offline-1",
                "installmentId": SAMPLE_INSTALLMENT_ID,
                "amountPaid": 1000000,
                "mode": "cash",
                "utrRef": ""
            },
            {
                "localId": "offline-2",
                "installmentId": "00000000-0000-0000-0000-000000000000",
                "amountPaid": 500000,
                "mode": "upi",
                "utrRef": "UPI123"
            }
        ], indent=2),
        "status_code": 200,
        "response": json.dumps([
            {
                "localId": "offline-1",
                "success": True,
                "error": "",
                "receiptId": "receipt-guid-1"
            },
            {
                "localId": "offline-2",
                "success": False,
                "error": "Installment not found",
                "receiptId": None
            }
        ], indent=2)
    },
    
    # PARTNERS CONTROLLER
    {
        "id": "TS_API_066",
        "endpoint": "GET /api/v1/partners",
        "description": "Verify whether the get all partners endpoint returns 200 response with list of partners",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_PARTNER_ID,
                "name": "Partner Company Ltd",
                "email": "partner@company.com",
                "phone": "9876543210",
                "isActive": True
            }
        ], indent=2)
    },
    {
        "id": "TS_API_067",
        "endpoint": f"GET /api/v1/partners/{SAMPLE_PARTNER_ID}",
        "description": "Verify whether the get partner by ID endpoint returns 200 response with partner details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "id": SAMPLE_PARTNER_ID,
            "name": "Partner Company Ltd",
            "email": "partner@company.com",
            "phone": "9876543210",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_068",
        "endpoint": "GET /api/v1/partners/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get partner by ID endpoint returns 404 response when partner does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Partner not found"
        }, indent=2)
    },
    {
        "id": "TS_API_069",
        "endpoint": "POST /api/v1/partners",
        "description": "Verify whether the create partner endpoint returns 201 response with valid partner data",
        "payload": json.dumps({
            "name": "New Partner Ltd",
            "email": "newpartner@company.com",
            "phone": "9123456789",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-partner-guid",
            "name": "New Partner Ltd",
            "email": "newpartner@company.com",
            "phone": "9123456789",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_070",
        "endpoint": "POST /api/v1/partners",
        "description": "Verify whether the create partner endpoint returns 400 response when name is missing",
        "payload": json.dumps({
            "email": "newpartner@company.com",
            "phone": "9123456789",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Name": ["The Name field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_071",
        "endpoint": "POST /api/v1/partners",
        "description": "Verify whether the create partner endpoint returns 400 response when email format is invalid",
        "payload": json.dumps({
            "name": "New Partner Ltd",
            "email": "invalid-email",
            "phone": "9123456789",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Email": ["The Email field is not a valid e-mail address."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_072",
        "endpoint": f"PUT /api/v1/partners/{SAMPLE_PARTNER_ID}",
        "description": "Verify whether the update partner endpoint returns 204 response with valid partner data",
        "payload": json.dumps({
            "name": "Updated Partner Ltd",
            "email": "partner@company.com",
            "phone": "9876543210",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_073",
        "endpoint": "PUT /api/v1/partners/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the update partner endpoint returns 404 response when partner does not exist",
        "payload": json.dumps({
            "name": "Updated Partner Ltd",
            "email": "partner@company.com",
            "phone": "9876543210",
            "branchId": SAMPLE_BRANCH_ID
        }, indent=2),
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Partner not found"
        }, indent=2)
    },
    
    # CAPITAL ACCOUNTS CONTROLLER
    {
        "id": "TS_API_074",
        "endpoint": "POST /api/v1/capitalaccounts/investment",
        "description": "Verify whether the add investment endpoint returns 200 response with valid investment data",
        "payload": json.dumps({
            "partnerId": SAMPLE_PARTNER_ID,
            "amount": 50000000
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "message": "Investment added successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_075",
        "endpoint": "POST /api/v1/capitalaccounts/investment",
        "description": "Verify whether the add investment endpoint returns 400 response when partnerId is missing",
        "payload": json.dumps({
            "amount": 50000000
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "PartnerId": ["The PartnerId field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_076",
        "endpoint": "POST /api/v1/capitalaccounts/investment",
        "description": "Verify whether the add investment endpoint returns 400 response when amount is missing",
        "payload": json.dumps({
            "partnerId": SAMPLE_PARTNER_ID
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Amount": ["The Amount field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_077",
        "endpoint": "POST /api/v1/capitalaccounts/investment",
        "description": "Verify whether the add investment endpoint returns 404 response when partner does not exist",
        "payload": json.dumps({
            "partnerId": "00000000-0000-0000-0000-000000000000",
            "amount": 50000000
        }, indent=2),
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Partner not found"
        }, indent=2)
    },
    {
        "id": "TS_API_078",
        "endpoint": "POST /api/v1/capitalaccounts/withdrawal",
        "description": "Verify whether the withdrawal endpoint returns 200 response with valid withdrawal data",
        "payload": json.dumps({
            "partnerId": SAMPLE_PARTNER_ID,
            "amount": 10000000
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "message": "Withdrawal recorded successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_079",
        "endpoint": "POST /api/v1/capitalaccounts/withdrawal",
        "description": "Verify whether the withdrawal endpoint returns 400 response when partnerId is missing",
        "payload": json.dumps({
            "amount": 10000000
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "PartnerId": ["The PartnerId field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_080",
        "endpoint": "POST /api/v1/capitalaccounts/withdrawal",
        "description": "Verify whether the withdrawal endpoint returns 400 response when withdrawal amount exceeds available balance",
        "payload": json.dumps({
            "partnerId": SAMPLE_PARTNER_ID,
            "amount": 999999999999
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "Insufficient balance"
        }, indent=2)
    },
    {
        "id": "TS_API_081",
        "endpoint": f"GET /api/v1/capitalaccounts/summary/{SAMPLE_PARTNER_ID}",
        "description": "Verify whether the get capital summary endpoint returns 200 response with partner capital details",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "partnerId": SAMPLE_PARTNER_ID,
            "partnerName": "Partner Company Ltd",
            "totalInvestment": 50000000,
            "totalProfit": 5000000,
            "totalWithdrawal": 10000000,
            "currentBalance": 45000000
        }, indent=2)
    },
    {
        "id": "TS_API_082",
        "endpoint": "GET /api/v1/capitalaccounts/summary/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get capital summary endpoint returns 404 response when partner does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Partner not found"
        }, indent=2)
    },

    # PRODUCT CONTROLLER
    {
        "id": "TS_API_083",
        "endpoint": "GET /api/v1/product/active",
        "description": "Verify whether the get active products endpoint returns 200 response with list of active loan products",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": SAMPLE_PRODUCT_ID,
                "branchId": SAMPLE_BRANCH_ID,
                "name": "Personal Loan",
                "code": "PL001",
                "interestRate": 12.5,
                "defaultTenureMonths": 12,
                "repaymentFrequency": "monthly",
                "isActive": True
            }
        ], indent=2)
    },
    {
        "id": "TS_API_084",
        "endpoint": "POST /api/v1/product",
        "description": "Verify whether the create product endpoint returns 201 response with valid product data",
        "payload": json.dumps({
            "branchId": SAMPLE_BRANCH_ID,
            "name": "Business Loan",
            "code": "BL001",
            "interestRate": 15.0,
            "defaultTenureMonths": 24,
            "repaymentFrequency": 0,
            "isActive": True
        }, indent=2),
        "status_code": 201,
        "response": json.dumps({
            "id": "new-product-guid",
            "branchId": SAMPLE_BRANCH_ID,
            "name": "Business Loan",
            "code": "BL001",
            "interestRate": 15.0,
            "defaultTenureMonths": 24,
            "repaymentFrequency": "monthly",
            "isActive": True
        }, indent=2)
    },
    {
        "id": "TS_API_085",
        "endpoint": "POST /api/v1/product",
        "description": "Verify whether the create product endpoint returns 400 response when name is missing",
        "payload": json.dumps({
            "branchId": SAMPLE_BRANCH_ID,
            "code": "BL001",
            "interestRate": 15.0,
            "defaultTenureMonths": 24,
            "repaymentFrequency": 0,
            "isActive": True
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "Name": ["The Name field is required."]
            }
        }, indent=2)
    },
    {
        "id": "TS_API_086",
        "endpoint": "POST /api/v1/product",
        "description": "Verify whether the create product endpoint returns 403 response when user is not super_admin or branch_manager",
        "payload": json.dumps({
            "branchId": SAMPLE_BRANCH_ID,
            "name": "Business Loan",
            "code": "BL001",
            "interestRate": 15.0,
            "defaultTenureMonths": 24,
            "repaymentFrequency": 0,
            "isActive": True
        }, indent=2),
        "status_code": 403,
        "response": json.dumps({
            "statusCode": 403,
            "message": "Insufficient permissions"
        }, indent=2)
    },
    {
        "id": "TS_API_087",
        "endpoint": f"DELETE /api/v1/product/{SAMPLE_PRODUCT_ID}",
        "description": "Verify whether the deactivate product endpoint returns 204 response when product is successfully deactivated",
        "payload": "N/A",
        "status_code": 204,
        "response": "No Content"
    },
    {
        "id": "TS_API_088",
        "endpoint": "DELETE /api/v1/product/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the deactivate product endpoint returns 404 response when product does not exist",
        "payload": "N/A",
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Product not found"
        }, indent=2)
    },
    
    # RECEIPTS CONTROLLER
    {
        "id": "TS_API_089",
        "endpoint": "POST /api/v1/receipts",
        "description": "Verify whether the record payment endpoint returns 200 response with valid payment data",
        "payload": json.dumps({
            "installmentId": SAMPLE_INSTALLMENT_ID,
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 200,
        "response": json.dumps({
            "id": "receipt-guid",
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": "",
            "capturedAt": "2026-04-11T10:30:00Z"
        }, indent=2)
    },
    {
        "id": "TS_API_090",
        "endpoint": "POST /api/v1/receipts",
        "description": "Verify whether the record payment endpoint returns 400 response when installmentId is missing",
        "payload": json.dumps({
            "amountPaid": 1000000,
            "mode": "cash",
            "utrRef": ""
        }, indent=2),
        "status_code": 400,
        "response": json.dumps({
            "errors": {
                "InstallmentId": ["The InstallmentId field is required."]
            }
        }, indent=2)
    },
    
    # REPORT CONTROLLER
    {
        "id": "TS_API_091",
        "endpoint": "GET /api/v1/report/par?start=2026-01-01&end=2026-04-11",
        "description": "Verify whether the get PAR report endpoint returns 200 response with portfolio at risk data",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "par30": 5.2,
            "par60": 2.1,
            "par90": 0.8,
            "totalOutstanding": 100000000,
            "overdueAmount": 5200000
        }, indent=2)
    },
    {
        "id": "TS_API_092",
        "endpoint": "GET /api/v1/report/par",
        "description": "Verify whether the get PAR report endpoint returns 200 response without date parameters",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "par30": 5.2,
            "par60": 2.1,
            "par90": 0.8,
            "totalOutstanding": 100000000,
            "overdueAmount": 5200000
        }, indent=2)
    },
    {
        "id": "TS_API_093",
        "endpoint": "GET /api/v1/report/efficiency?start=2026-01-01&end=2026-04-11",
        "description": "Verify whether the get collection efficiency report endpoint returns 200 response with efficiency data",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "collectionRate": 95.5,
            "totalDue": 50000000,
            "totalCollected": 47750000,
            "pendingAmount": 2250000
        }, indent=2)
    },
    {
        "id": "TS_API_094",
        "endpoint": "GET /api/v1/report/efficiency",
        "description": "Verify whether the get collection efficiency report endpoint returns 200 response without date parameters",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "collectionRate": 95.5,
            "totalDue": 50000000,
            "totalCollected": 47750000,
            "pendingAmount": 2250000
        }, indent=2)
    },
    {
        "id": "TS_API_095",
        "endpoint": "GET /api/v1/report/dashboard-stats",
        "description": "Verify whether the get dashboard stats endpoint returns 200 response with dashboard statistics",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "totalLoans": 150,
            "activeLoans": 120,
            "totalDisbursed": 150000000,
            "totalCollected": 95000000,
            "totalOutstanding": 55000000,
            "overdueLoans": 8
        }, indent=2)
    },
    {
        "id": "TS_API_096",
        "endpoint": "GET /api/v1/report/par",
        "description": "Verify whether the get PAR report endpoint returns 403 response when user is not authorized",
        "payload": "N/A",
        "status_code": 403,
        "response": json.dumps({
            "statusCode": 403,
            "message": "Insufficient permissions"
        }, indent=2)
    },
    
    # AUDIT CONTROLLER
    {
        "id": "TS_API_097",
        "endpoint": f"GET /api/v1/audit/LoanCase/{SAMPLE_LOAN_ID}",
        "description": "Verify whether the get entity history endpoint returns 200 response with audit logs for a loan case",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": "audit-log-guid-1",
                "entityName": "LoanCase",
                "recordId": SAMPLE_LOAN_ID,
                "action": "Created",
                "userId": SAMPLE_USER_ID,
                "userName": "John Doe",
                "timestamp": "2026-04-01T10:00:00Z",
                "changes": "{\"Status\":\"draft\"}"
            },
            {
                "id": "audit-log-guid-2",
                "entityName": "LoanCase",
                "recordId": SAMPLE_LOAN_ID,
                "action": "Updated",
                "userId": SAMPLE_USER_ID,
                "userName": "John Doe",
                "timestamp": "2026-04-05T14:30:00Z",
                "changes": "{\"Status\":\"active\"}"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_098",
        "endpoint": "GET /api/v1/audit/Customer/00000000-0000-0000-0000-000000000000",
        "description": "Verify whether the get entity history endpoint returns 200 response with empty array when no audit logs exist",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([], indent=2)
    },
    {
        "id": "TS_API_099",
        "endpoint": "GET /api/v1/audit/recent?count=50",
        "description": "Verify whether the get recent logs endpoint returns 200 response with recent audit logs",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": "audit-log-guid",
                "entityName": "LoanCase",
                "recordId": SAMPLE_LOAN_ID,
                "action": "Updated",
                "userId": SAMPLE_USER_ID,
                "userName": "John Doe",
                "timestamp": "2026-04-11T09:00:00Z",
                "changes": "{\"Status\":\"active\"}"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_100",
        "endpoint": "GET /api/v1/audit/recent",
        "description": "Verify whether the get recent logs endpoint returns 200 response with default count of 100 logs",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": "audit-log-guid",
                "entityName": "LoanCase",
                "recordId": SAMPLE_LOAN_ID,
                "action": "Updated",
                "userId": SAMPLE_USER_ID,
                "userName": "John Doe",
                "timestamp": "2026-04-11T09:00:00Z",
                "changes": "{}"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_101",
        "endpoint": "GET /api/v1/audit/recent",
        "description": "Verify whether the get recent logs endpoint returns 403 response when user is not super_admin or accountant",
        "payload": "N/A",
        "status_code": 403,
        "response": json.dumps({
            "statusCode": 403,
            "message": "Insufficient permissions"
        }, indent=2)
    },
    
    # RECOVERY CONTROLLER
    {
        "id": "TS_API_102",
        "endpoint": "GET /api/recovery/overdue",
        "description": "Verify whether the get overdue loans endpoint returns 200 response with list of overdue loans",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "loanId": SAMPLE_LOAN_ID,
                "customerId": SAMPLE_CUSTOMER_ID,
                "customerName": "Rajesh Kumar",
                "principal": 10000000,
                "overdueAmount": 1500000,
                "daysPastDue": 15,
                "lastPaymentDate": "2026-03-20T00:00:00Z"
            }
        ], indent=2)
    },
    {
        "id": "TS_API_103",
        "endpoint": f"POST /api/recovery/{SAMPLE_LOAN_ID}/follow-up",
        "description": "Verify whether the record follow-up endpoint returns 200 response with valid follow-up notes",
        "payload": json.dumps("Customer contacted via phone. Promised to pay by end of week.", indent=2),
        "status_code": 200,
        "response": json.dumps({
            "message": "Follow-up recorded"
        }, indent=2)
    },
    {
        "id": "TS_API_104",
        "endpoint": "POST /api/recovery/00000000-0000-0000-0000-000000000000/follow-up",
        "description": "Verify whether the record follow-up endpoint returns 404 response when loan does not exist",
        "payload": json.dumps("Follow-up notes", indent=2),
        "status_code": 404,
        "response": json.dumps({
            "statusCode": 404,
            "message": "Loan case not found"
        }, indent=2)
    },
    
    # DAY END CONTROLLER
    {
        "id": "TS_API_105",
        "endpoint": "POST /api/dayend/close?date=2026-04-10&verifiedCash=5000000",
        "description": "Verify whether the close day end endpoint returns 200 response with valid date and cash amount",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "message": "Day 2026-04-10 closed successfully"
        }, indent=2)
    },
    {
        "id": "TS_API_106",
        "endpoint": "POST /api/dayend/close?date=2026-04-10&verifiedCash=-1000",
        "description": "Verify whether the close day end endpoint returns 400 response when verified cash is negative",
        "payload": "N/A",
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "Verified cash cannot be negative"
        }, indent=2)
    },
    {
        "id": "TS_API_107",
        "endpoint": "POST /api/dayend/close?date=2026-04-10&verifiedCash=5000000",
        "description": "Verify whether the close day end endpoint returns 400 response when day is already closed",
        "payload": "N/A",
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "Day already closed"
        }, indent=2)
    },
    
    # JOURNAL CONTROLLER
    {
        "id": "TS_API_108",
        "endpoint": "GET /api/journal/entries",
        "description": "Verify whether the get journal entries endpoint returns 200 response with list of journal entries",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps([
            {
                "id": "journal-entry-guid",
                "publicId": "JE-2026-001",
                "date": "2026-04-10T00:00:00Z",
                "description": "Loan disbursement",
                "reference": SAMPLE_LOAN_ID,
                "lines": [
                    {
                        "accountName": "Loans Receivable",
                        "type": "debit",
                        "amount": 10000000
                    },
                    {
                        "accountName": "Cash",
                        "type": "credit",
                        "amount": 10000000
                    }
                ]
            }
        ], indent=2)
    },
    
    # LEDGER CONTROLLER
    {
        "id": "TS_API_109",
        "endpoint": "GET /api/ledger/trial-balance",
        "description": "Verify whether the get trial balance endpoint returns 200 response with trial balance data",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "accounts": [
                {
                    "accountName": "Cash",
                    "debit": 50000000,
                    "credit": 0
                },
                {
                    "accountName": "Loans Receivable",
                    "debit": 100000000,
                    "credit": 0
                },
                {
                    "accountName": "Capital",
                    "debit": 0,
                    "credit": 150000000
                }
            ],
            "totalDebit": 150000000,
            "totalCredit": 150000000
        }, indent=2)
    },
    {
        "id": "TS_API_110",
        "endpoint": "GET /api/ledger/pnl?start=2026-01-01&end=2026-04-11",
        "description": "Verify whether the get P&L endpoint returns 200 response with profit and loss statement",
        "payload": "N/A",
        "status_code": 200,
        "response": json.dumps({
            "profitAndLoss": {
                "revenue": {
                    "interestIncome": 5000000,
                    "processingFees": 500000,
                    "total": 5500000
                },
                "expenses": {
                    "operatingExpenses": 1000000,
                    "salaries": 2000000,
                    "total": 3000000
                },
                "netProfit": 2500000
            }
        }, indent=2)
    },
    {
        "id": "TS_API_111",
        "endpoint": "GET /api/ledger/pnl?start=2026-04-11&end=2026-01-01",
        "description": "Verify whether the get P&L endpoint returns 400 response when start date is after end date",
        "payload": "N/A",
        "status_code": 400,
        "response": json.dumps({
            "statusCode": 400,
            "message": "Invalid date range"
        }, indent=2)
    },
]

# Write test cases to Excel
row_num = 2
for tc in test_cases:
    ws.cell(row=row_num, column=1).value = tc["id"]
    ws.cell(row=row_num, column=2).value = tc["endpoint"]
    ws.cell(row=row_num, column=3).value = tc["description"]
    ws.cell(row=row_num, column=4).value = tc["payload"]
    ws.cell(row=row_num, column=5).value = tc["status_code"]
    ws.cell(row=row_num, column=6).value = tc["response"]
    ws.cell(row=row_num, column=7).value = ""  # Status column
    ws.cell(row=row_num, column=8).value = ""  # Comments column
    
    # Apply borders and alignment
    for col_num in range(1, 9):
        cell = ws.cell(row=row_num, column=col_num)
        cell.border = border
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    
    row_num += 1

# Adjust column widths
ws.column_dimensions['A'].width = 15  # Test Case ID
ws.column_dimensions['B'].width = 40  # Endpoint Name
ws.column_dimensions['C'].width = 60  # Test Case Description
ws.column_dimensions['D'].width = 50  # Request Payload
ws.column_dimensions['E'].width = 20  # Expected Status Code
ws.column_dimensions['F'].width = 50  # Expected Response
ws.column_dimensions['G'].width = 15  # Status
ws.column_dimensions['H'].width = 30  # Comments

# Freeze the header row
ws.freeze_panes = 'A2'

# Save the workbook
wb.save('Fintech_API_Test_Cases.xlsx')
print("Excel file 'Fintech_API_Test_Cases.xlsx' has been created successfully!")
print(f"Total test cases: {len(test_cases)}")
