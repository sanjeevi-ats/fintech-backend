import openpyxl
import requests
import json
import re

TEST_EXCEL = 'FinVeda_API_Test_Cases.xlsx'
OUTPUT_EXCEL = 'FinVeda_Test_Execution_Report.xlsx'
BASE_URL = 'http://localhost:5177'

def infer_method(desc, endpoint, payload):
    desc_lower = desc.lower()
    if 'delete ' in desc_lower: return 'DELETE'
    if 'update ' in desc_lower or 'modify ' in desc_lower: return 'PUT'
    
    if 'search' in endpoint.lower(): return 'GET'
    if '/status' in endpoint.lower(): return 'GET'
    if '/dcs' in endpoint.lower(): return 'GET'
    if '/logs' in endpoint.lower(): return 'GET'
    
    if '/update' in endpoint.lower(): return 'PUT'
    
    if 'approve' in endpoint.lower() or 'disburse' in endpoint.lower() or 'close' in endpoint.lower() or 'waive' in endpoint.lower() or 'collect' in endpoint.lower() or 'distribute' in endpoint.lower() or 'apply' in endpoint.lower():
        if 'delete' not in desc_lower and 'update' not in desc_lower:
            return 'POST'
    
    if 'create ' in desc_lower or 'register ' in desc_lower or 'login ' in desc_lower or 'add ' in desc_lower or 'assign ' in desc_lower: return 'POST'
    if 'get ' in desc_lower or 'returns list ' in desc_lower or 'returns data ' in desc_lower: return 'GET'
    if 'returns 400' in desc_lower and payload != 'N/A' and payload:
        if 'login' in endpoint.lower() or 'register' in endpoint.lower() or 'create' in endpoint.lower():
            return 'POST'
        if 'update' in endpoint.lower():
            return 'PUT'
        
    if 'returns 401' in desc_lower or 'returns 403' in desc_lower or 'returns 404' in desc_lower:
        if 'login' in endpoint.lower() or 'register' in endpoint.lower() or 'create' in endpoint.lower():
            return 'POST'
        if 'update' in endpoint.lower():
            return 'PUT'
        if 'delete' in endpoint.lower():
            return 'DELETE'
        return 'GET'
        
    if payload and payload != 'N/A':
        return 'POST'
    return 'GET'

def parse_payload(payload_str):
    if not payload_str or payload_str == 'N/A':
        return None
    # Remove context strings like 'N/A (Query params: ...)' or 'N/A (Path: ...)'
    if 'N/A' in payload_str and '(' in payload_str:
        return None
        
    try:
        return json.loads(payload_str)
    except json.JSONDecodeError:
        return None

def main():
    wb = openpyxl.load_workbook(TEST_EXCEL)
    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active) # Remove default sheet
    
    auth_token = None
    pass_count = 0
    fail_count = 0
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        out_ws = out_wb.create_sheet(sheet_name)
        
        # Copy headers and add Actual Status Code, Actual Response, Final Status
        headers = [str(cell.value) if cell.value else '' for cell in ws[1]]
        out_ws.append(headers + ['Actual Status Code', 'Actual Response', 'Execution Status', 'Execution Comments'])
        
        for row_idx in range(2, ws.max_row + 1):
            row_vals = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
            row_strs = [str(v) if v is not None else '' for v in row_vals]
            
            tc_id, endpoint, desc, payload_raw, exp_status_str, exp_response_raw, status, comments = row_strs[:8]
            
            if not tc_id or tc_id == 'N/A':
                continue
                
            method = infer_method(desc, endpoint, payload_raw)
            url = endpoint
            payload = parse_payload(payload_raw)
            
            try:
                exp_status = int(float(exp_status_str))
            except ValueError:
                exp_status = 0
                
            headers_req = {'Content-Type': 'application/json'}
            if auth_token and ('401' not in desc and 'unauthorized' not in desc.lower()):
                headers_req['Authorization'] = f'Bearer {auth_token}'
                
            if 'admin token' in desc.lower() or 'authorized request is made' in desc.lower() or 'success' in desc.lower():
                 if not auth_token: # Need to login first to get a token, but let's just use what we have, or login explicitly
                      pass
                      
            print(f"Executing {tc_id}: {method} {url}")
            
            actual_status = 0
            actual_response_text = ""
            try:
                if method == 'GET':
                    res = requests.get(url, headers=headers_req, timeout=5)
                elif method == 'POST':
                    res = requests.post(url, headers=headers_req, json=payload, timeout=5)
                elif method == 'PUT':
                    res = requests.put(url, headers=headers_req, json=payload, timeout=5)
                elif method == 'DELETE':
                    res = requests.delete(url, headers=headers_req, timeout=5)
                else:
                    res = requests.get(url, headers=headers_req, timeout=5)
                    
                actual_status = res.status_code
                try:
                    actual_response_text = json.dumps(res.json(), indent=2)
                except:
                    actual_response_text = res.text
                    
                # If Login success, capture token
                if 'login' in url.lower() and actual_status == 200:
                    try:
                        auth_token = res.json().get('token', auth_token)
                    except:
                        pass
            except Exception as e:
                actual_status = 500
                actual_response_text = str(e)
                
            # Evaluation
            exec_status = "PASSED" if actual_status == exp_status else "FAILED"
            
            if exec_status == "PASSED":
                pass_count += 1
            else:
                fail_count += 1
                
            exec_comments = f"Expected {exp_status}, got {actual_status}. Method inferred: {method}"
            
            out_row = list(row_vals) + [actual_status, actual_response_text, exec_status, exec_comments]
            out_ws.append(out_row)
            
    out_wb.save(OUTPUT_EXCEL)
    print(f"\\nExecution complete! Passed: {pass_count}, Failed: {fail_count}")

if __name__ == "__main__":
    main()
