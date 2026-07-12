import openpyxl

wb = openpyxl.load_workbook('FinVeda_API_Test_Cases.xlsx')
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"--- {sheet} ---")
    headers = [str(cell.value) if cell.value else '' for cell in ws[1]]
    print("HEADERS:", headers)
    for row_idx in range(2, min(5, ws.max_row + 1)):
        row = [str(ws.cell(row=row_idx, column=c).value) if ws.cell(row=row_idx, column=c).value else '' for c in range(1, len(headers) + 1)]
        print(f"ROW {row_idx}: {row}")
    break
