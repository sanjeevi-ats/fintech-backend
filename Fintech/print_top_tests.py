import openpyxl

wb = openpyxl.load_workbook('FinVeda_Test_Execution_Report.xlsx')
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(ws.cell(row=12+1, column=1).value, ws.cell(row=12+1, column=2).value, ws.cell(row=12+1, column=11).value)
    
